"""Мини-библиотека для презентаций 16:9 на python-pptx с темами оформления.

Размеры в дюймах, слайд 13.333 × 7.5. Три темы: NOIR (чёрный и красный, Arial Black),
EDITORIAL (кремовый фон, терракота, Georgia) и FRESH (светлый, зелёный, Trebuchet MS, скругления).
Все шрифты есть в Windows и macOS.
"""
from dataclasses import dataclass

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = 13.333, 7.5
CONCEPT = "Концепт для портфолио, бренд и данные вымышлены"
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


@dataclass(frozen=True)
class Theme:
    name: str
    bg: RGBColor
    panel: RGBColor
    panel2: RGBColor
    acc: RGBColor
    acc2: RGBColor
    text: RGBColor
    muted: RGBColor
    dim: RGBColor
    on_acc: RGBColor
    acc_ink: RGBColor
    head: str
    body: str
    head_bold: bool
    rounded: bool
    card_accent: str
    mono: str = "Courier New"


NOIR = Theme("noir", rgb("0B0B0C"), rgb("17171A"), rgb("222226"), rgb("E3202B"), rgb("5A0D12"),
             rgb("FFFFFF"), rgb("A3A3A8"), rgb("5C5C63"), rgb("FFFFFF"), rgb("0B0B0C"),
             "Arial Black", "Arial", False, False, "top")
EDITORIAL = Theme("editorial", rgb("F3EEE4"), rgb("E9E1D3"), rgb("DDD2BF"), rgb("B9502C"), rgb("6E3520"),
                  rgb("1F1B17"), rgb("5E554B"), rgb("9A8E7E"), rgb("FFFFFF"), rgb("F3EEE4"),
                  "Georgia", "Arial", True, False, "left")
FRESH = Theme("fresh", rgb("F4F8F3"), rgb("E2EEE5"), rgb("CDE3D4"), rgb("1E7A4C"), rgb("0F4A2E"),
              rgb("13261C"), rgb("4D6255"), rgb("8CA295"), rgb("FFFFFF"), rgb("FFFFFF"),
              "Trebuchet MS", "Trebuchet MS", True, True, "none")


class Deck:
    def __init__(self, title: str, theme: Theme = NOIR):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.prs.core_properties.title = title
        self.title = title
        self.t = theme
        self.count = 0
        self._pending = None

    def slide(self, number: bool = True, footer_x: float = 0.6, footer_color=None):
        self._footer()
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = self.t.bg
        self.count += 1
        self.s = s
        self._pending = (number, footer_x, footer_color)
        return s

    def _footer(self):
        if not self._pending:
            return
        number, footer_x, color = self._pending
        self.text(footer_x, 7.02, 7.5, 0.3, CONCEPT, 10, color=color or self.t.muted)
        if number:
            self.text(11.9, 7.02, 0.83, 0.3, f"{self.count:02d}", 10, color=self.t.dim, align="r", font=self.t.head)
        self._pending = None

    def shape(self, x, y, w, h, fill=None, line=None, kind=MSO_SHAPE.RECTANGLE, line_w=1.5, rot=0, adj=None):
        sh = self.s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(line_w)
        if adj is not None:
            sh.adjustments[0] = adj
        sh.shadow.inherit = False
        if rot:
            sh.rotation = rot
        sh.text_frame.text = ""
        return sh

    def box(self, x, y, w, h, fill=None, line=None, line_w=1.5):
        if self.t.rounded:
            return self.shape(x, y, w, h, fill, line, MSO_SHAPE.ROUNDED_RECTANGLE, line_w,
                              adj=min(0.5, 0.18 / min(w, h)))
        return self.shape(x, y, w, h, fill, line, line_w=line_w)

    def text(self, x, y, w, h, content, size, *, color=None, font=None, bold=False, align="l", anchor="t",
             spacing=1.15, shape=None):
        color = color or self.t.text
        font = font or self.t.body
        if shape is None:
            shape = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = ANCHOR[anchor]
        paras = content.split("\n") if isinstance(content, str) else content
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            p.line_spacing = spacing
            runs = [(para, {})] if isinstance(para, str) else para
            for txt, opt in runs:
                r = p.add_run()
                r.text = txt
                f = r.font
                run_font = opt.get("font", font)
                f.size = Pt(opt.get("size", size))
                f.bold = opt.get("bold", bold or (run_font == self.t.head and self.t.head_bold))
                f.name = run_font
                f.color.rgb = opt.get("color", color)
        return shape

    def h(self, x, y, w, h, content, size, **kw):
        kw.setdefault("spacing", 1.0)
        return self.text(x, y, w, h, content, size, font=self.t.head, **kw)

    def kicker(self, x, y, label):
        if self.t.name == "editorial":
            self.shape(x, y + 0.32, 0.9, 0.02, fill=self.t.acc)
            self.text(x, y, 8, 0.3, label.upper(), 11, color=self.t.acc, bold=True)
        elif self.t.name == "fresh":
            sh = self.box(x, y - 0.04, 0.3 + 0.105 * len(label), 0.36, fill=self.t.panel2)
            self.text(0, 0, 0, 0, label, 12, color=self.t.acc2, bold=True, align="c", anchor="m", shape=sh)
        else:
            self.shape(x, y + 0.08, 0.35, 0.08, fill=self.t.acc)
            self.text(x + 0.5, y, 8, 0.3, label.upper(), 12, color=self.t.acc, bold=True)

    def heading(self, title, kicker=None, y=0.55, size=36, w=12.1):
        if kicker:
            self.kicker(0.6, y, kicker)
            y += 0.45
        self.h(0.6, y, w, 1.5, title, size)

    def badge(self, x, y, d, label, fill=None, color=None, size=20, line=None):
        fill = self.t.acc if fill is None and line is None else fill
        sh = self.shape(x, y, d, d, fill=fill, line=line, kind=MSO_SHAPE.OVAL)
        self.h(0, 0, 0, 0, label, size, color=color or self.t.on_acc, align="c", anchor="m", shape=sh)
        return sh

    def card(self, x, y, w, h, title, body, *, num=None, fill=None, title_size=18, body_size=14,
             title_h=0.75):
        self.box(x, y, w, h, fill=fill or self.t.panel)
        if self.t.card_accent == "top":
            self.shape(x, y, w, 0.07, fill=self.t.acc)
        elif self.t.card_accent == "left":
            self.shape(x, y, 0.05, h, fill=self.t.acc)
        ty = y + 0.3
        if num is not None:
            self.badge(x + 0.3, ty, 0.55, str(num), size=16)
            ty += 0.8
        self.h(x + 0.3, ty, w - 0.6, title_h, title, title_size)
        by = ty + title_h + 0.1
        self.text(x + 0.3, by, w - 0.6, y + h - by - 0.25, body, body_size, color=self.t.muted)

    def prompt_box(self, x, y, w, h, prompt, size=12):
        self.box(x, y, w, h, fill=self.t.panel2, line=self.t.dim, line_w=0.75)
        self.shape(x, y, 0.06, h, fill=self.t.acc)
        self.text(x + 0.25, y + 0.15, 2.5, 0.25, "ПРОМПТ", 9, color=self.t.acc, bold=True)
        self.text(x + 0.25, y + 0.45, w - 0.45, h - 0.6, prompt, size, font=self.t.mono, spacing=1.1)

    def save(self, path):
        self._footer()
        self.prs.save(path)
